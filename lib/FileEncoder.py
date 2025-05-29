import pathlib
import os
from docx import Document as DocxDocument
import json
from bs4 import BeautifulSoup
import csv
from reportlab.platypus import SimpleDocTemplate, Preformatted
from reportlab.lib.styles import getSampleStyleSheet
from openpyxl import Workbook

def write_text_to_txt(file_path: str, content: str) -> str:
    """Writes plain text content to a .txt file."""
    try:
        path = pathlib.Path(file_path).resolve()
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(content, encoding='utf-8')
        return f"Successfully wrote text content to '{path}'."
    except Exception as e:
        return f"Error writing text to '{file_path}': {e}"

def write_text_to_docx(file_path: str, content: str) -> str:
    """Writes text content to a basic .docx file."""
    try:
        path = pathlib.Path(file_path).resolve()
        path.parent.mkdir(parents=True, exist_ok=True)
        document = DocxDocument()
        # Add content line by line as paragraphs
        for line in content.splitlines():
            document.add_paragraph(line)
        document.save(path)
        return f"Successfully wrote DOCX content to '{path}'."
    except Exception as e:
        return f"Error writing DOCX to '{file_path}': {e}"

def write_text_to_html(file_path: str, content: str) -> str:
    """Writes text content to a basic .html file."""
    try:
        path = pathlib.Path(file_path).resolve()
        path.parent.mkdir(parents=True, exist_ok=True)
        # Simple HTML structure
        html_content = f"""<!DOCTYPE html>
<html>
<head>
<title>Document</title>
</head>
<body>
<pre>{content}</pre>
</body>
</html>"""
        path.write_text(html_content, encoding='utf-8')
        return f"Successfully wrote HTML content to '{path}'."
    except Exception as e:
        return f"Error writing HTML to '{file_path}': {e}"

def write_text_to_json(file_path: str, content: str) -> str:
    """Attempts to parse content as JSON and writes it to a .json file."""
    try:
        path = pathlib.Path(file_path).resolve()
        path.parent.mkdir(parents=True, exist_ok=True)
        # Attempt to parse content as JSON
        json_data = json.loads(content)
        with open(path, 'w', encoding='utf-8') as f:
            json.dump(json_data, f, indent=4)
        return f"Successfully wrote JSON content to '{path}'."
    except json.JSONDecodeError:
        return f"Error writing JSON to '{file_path}': Content is not valid JSON."
    except Exception as e:
        return f"Error writing JSON to '{file_path}': {e}"

# --- New function for PDF writing ---
def write_text_to_pdf(file_path: str, content: str) -> str:
    """Writes text content to a basic .pdf file, preserving line breaks and basic spacing."""
    try:
        path = pathlib.Path(file_path).resolve()
        path.parent.mkdir(parents=True, exist_ok=True)

        doc = SimpleDocTemplate(str(path))
        styles = getSampleStyleSheet()
        story = []

        # Use Preformatted for better preservation of text structure,
        # similar to how <pre> works in HTML.
        # The 'Code' style typically uses a monospaced font.
        preformatted_text = Preformatted(content, styles['Code'])
        story.append(preformatted_text)

        doc.build(story)
        return f"Successfully wrote PDF content to '{path}'."
    except Exception as e:
        return f"Error writing PDF to '{file_path}': {e}"
# --- End of new PDF function ---

# --- New function for Excel (.xlsx) writing ---
def write_text_to_xlsx(file_path: str, content: str) -> str:
    """Writes text content to an .xlsx file, with each line in a new row of the first column."""
    try:
        path = pathlib.Path(file_path).resolve()
        path.parent.mkdir(parents=True, exist_ok=True)

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Content"

        # Split content by lines and write each line to a new row in the first column
        for r_idx, line in enumerate(content.splitlines(), start=1):
            sheet.cell(row=r_idx, column=1, value=line)

        workbook.save(filename=str(path))
        return f"Successfully wrote XLSX content to '{path}'."
    except Exception as e:
        return f"Error writing XLSX to '{file_path}': {e}"
# --- End of new Excel function ---

def write_file_content(file_path_str: str, content: str) -> str:
    """
    Writes content to a file, determining the format based on the file extension.
    Returns a success message or an error message.
    """
    file_path = pathlib.Path(file_path_str)
    ext = file_path.suffix.lower()

    if ext == ".txt":
        return write_text_to_txt(file_path_str, content)
    elif ext == ".docx":
        return write_text_to_docx(file_path_str, content)
    elif ext in [".html", ".htm"]:
        return write_text_to_html(file_path_str, content)
    elif ext == ".json":
        return write_text_to_json(file_path_str, content)
    # --- Updated to include PDF and XLSX ---
    elif ext == ".pdf":
        return write_text_to_pdf(file_path_str, content)
    elif ext == ".xlsx":
        return write_text_to_xlsx(file_path_str, content)
    # --- End of update ---
    # Add more formats here as needed
    else:
        return (f"Error: Unsupported file extension for writing: '{ext}'. "
                f"Supported formats: .txt, .docx, .html, .json, .pdf, .xlsx.")

if __name__ == '__main__':
    # Create a 'test_outputs' directory for generated files
    output_dir = pathlib.Path("test_outputs_writer")
    output_dir.mkdir(exist_ok=True)

    sample_text_content = """Hello World!
This is a test document.
It contains multiple lines of text.
    Including some with leading spaces.
And special characters: áéíóúñ & < > " ' / \\.
1. First item
2. Second item
"""

    sample_json_content = """{
    "name": "Test Document",
    "version": 1.0,
    "items": [
        {"id": 1, "value": "Apple"},
        {"id": 2, "value": "Banana"}
    ]
}"""

    invalid_json_content = "This is not valid JSON { name: Test"

    print("--- Testing file writing ---")

    # Test TXT
    txt_file = output_dir / "sample_text.txt"
    print(write_file_content(str(txt_file), sample_text_content))

    # Test DOCX
    docx_file = output_dir / "sample_text.docx"
    print(write_file_content(str(docx_file), sample_text_content))

    # Test HTML
    html_file = output_dir / "sample_text.html"
    print(write_file_content(str(html_file), sample_text_content))

    # Test JSON (valid)
    json_file_valid = output_dir / "sample_data.json"
    print(write_file_content(str(json_file_valid), sample_json_content))

    # Test JSON (invalid)
    json_file_invalid = output_dir / "invalid_data.json"
    print(write_file_content(str(json_file_invalid), invalid_json_content)) # Expected JSON error

    # Test PDF
    pdf_file = output_dir / "sample_text.pdf"
    print(write_file_content(str(pdf_file), sample_text_content))

    # Test XLSX
    xlsx_file = output_dir / "sample_text.xlsx"
    print(write_file_content(str(xlsx_file), sample_text_content))

    # Test unsupported extension
    unsupported_file = output_dir / "sample_text.rtf"
    print(write_file_content(str(unsupported_file), sample_text_content))

    print(f"\nTest files generated in '{output_dir.resolve()}'")